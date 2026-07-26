#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Testes do Aviso de Dispensa Completo.

Cobrem os 25 testes obrigatórios do escopo (item 16). Os que dependem das
minutas-mãe reais rodam sobre `05_minutas/`; os demais usam documentos
sintéticos ou a fixture de TR em `99_testes/aviso_completo/fixtures/`.

A montagem completa roda uma única vez, em fixture de escopo de módulo, porque
é o caminho caro (une, padroniza e valida cinco documentos).

Rodar:
    python -m pytest 99_testes/aviso_completo/ -v
"""
from __future__ import annotations

import json
import shutil
import zipfile
from pathlib import Path

import pytest

docx = pytest.importorskip(
    "docx", reason="python-docx não instalado (requirements-docx.txt)")
pytest.importorskip(
    "docxcompose", reason="docxcompose não instalado (requirements-docx.txt)")

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Cm

import extrair_dados_tr
import gerar_modelo_proposta
import gerar_pacote_publicacao
import localizar_componentes
import montar_aviso_completo
import numerar_anexos
import unir_docx
import validar_aviso_completo
from ocorrencias import BLOQUEANTE, Registro
from util_ooxml import iter_paragrafos_corpo, texto_paragrafo

RAIZ = Path(__file__).resolve().parents[2]
FIXTURES = Path(__file__).resolve().parent / "fixtures"
MANIFESTO = FIXTURES / "manifesto_aviso_completo.json"
MANIFESTO_CONTRATO = FIXTURES / "manifesto_com_contrato.json"
TR_FIXTURE = FIXTURES / "TR_FINAL.docx"

MINUTA_AVISO = localizar_componentes.MINUTA_AVISO
MINUTA_PROPOSTA = localizar_componentes.MINUTA_PROPOSTA
MINUTA_CONTRATO = (RAIZ / "05_minutas" / "CONTRATO_COMPRAS"
                   / "CONTRATO_COMPRAS_ENTREGA_FORN_CONTINUO_MINUTA_MAE.docx")


# --------------------------------------------------------------------------- #
# Utilidades
# --------------------------------------------------------------------------- #

def texto_do(caminho: Path) -> str:
    return "\n".join(texto_paragrafo(p)
                     for p in iter_paragrafos_corpo(Document(str(caminho))))


def tr_sintetico(destino: Path, itens: list[tuple[str, str, str, str]],
                 processo: str = "026/2026", pendente: bool = False) -> Path:
    """TR mínimo com quadro de itens, para os cenários que a fixture não cobre."""
    documento = Document()
    documento.add_paragraph("TERMO DE REFERÊNCIA")
    documento.add_paragraph(f"Solicitação nº {processo}")
    documento.add_paragraph("OBJETO: Aquisição de materiais diversos para a Câmara.")
    if pendente:
        documento.add_paragraph("Dotação orçamentária: [PREENCHER: dotação]")
    tabela = documento.add_table(rows=1, cols=4)
    for celula, titulo in zip(tabela.rows[0].cells,
                              ("Item", "Especificação", "Unidade", "Quantidade")):
        celula.text = titulo
    for item, descricao, unidade, quantidade in itens:
        linha = tabela.add_row()
        for celula, valor in zip(linha.cells, (item, descricao, unidade, quantidade)):
            celula.text = valor
    destino.parent.mkdir(parents=True, exist_ok=True)
    documento.save(str(destino))
    return destino


def manifesto_ajustado(destino: Path, base: Path = MANIFESTO, **mudancas) -> Path:
    dados = json.loads(base.read_text(encoding="utf-8"))
    for chave, valor in mudancas.items():
        bloco, _, campo = chave.partition("__")
        if campo:
            dados.setdefault(bloco, {})[campo] = valor
        else:
            dados[bloco] = valor
    destino.write_text(json.dumps(dados, ensure_ascii=False, indent=2),
                       encoding="utf-8")
    return destino


@pytest.fixture(scope="module")
def montagem_sem_contrato(tmp_path_factory):
    """Montagem completa, sem contrato — o caminho feliz do módulo."""
    saida = tmp_path_factory.mktemp("sem_contrato")
    dados = montar_aviso_completo.montar(
        None, MANIFESTO, saida=saida, converter_pdf=False)
    return dados, saida


# --------------------------------------------------------------------------- #
# 1 e 2 — com e sem contrato
# --------------------------------------------------------------------------- #

def test_01_aviso_com_contrato_numera_cinco_anexos():
    registro = Registro()
    manifesto = json.loads(MANIFESTO_CONTRATO.read_text(encoding="utf-8"))
    componentes = localizar_componentes.localizar(
        manifesto, registro, tr_explicito=TR_FIXTURE)
    plano = numerar_anexos.planejar(componentes, registro)

    assert [anexo.romano for anexo in plano.anexos] == ["I", "II", "III", "IV", "V"]
    assert plano.por_chave("contrato").romano == "IV"
    assert componentes.instrumento.tipo == "contrato"
    assert componentes.instrumento.minuta == MINUTA_CONTRATO


def test_01b_montagem_com_contrato_bloqueia_por_bloco_ou(tmp_path):
    """A minuta de contrato traz alternativas 'OU' que exigem decisão humana."""
    dados = montar_aviso_completo.montar(
        None, MANIFESTO_CONTRATO, saida=tmp_path, converter_pdf=False)
    mensagens = [o["mensagem"] for o in dados["ocorrencias"]
                 if o["severidade"] == BLOQUEANTE]
    assert any("'OU'" in m for m in mensagens)
    assert dados["status"] == "BLOQUEADO"


def test_02_aviso_sem_contrato_gera_documento(montagem_sem_contrato):
    dados, saida = montagem_sem_contrato
    final = saida / "saida" / "AVISO_DISPENSA_COMPLETO.docx"

    assert final.exists()
    assert dados["status"] in ("APTO PARA CONFERÊNCIA", "APTO COM RESSALVAS")
    assert dados["contagem"][BLOQUEANTE] == 0
    assert [a["romano"] for a in dados["anexos"]] == ["I", "II", "III", "IV"]


# --------------------------------------------------------------------------- #
# 3 e 4 — um item e vários itens
# --------------------------------------------------------------------------- #

def test_03_tr_com_um_item():
    dados = extrair_dados_tr.extrair(TR_FIXTURE, Registro())
    assert len(dados.itens) == 1
    assert dados.itens[0].quantidade == "1"
    assert "micro-ondas" in dados.itens[0].descricao.lower()


def test_04_tr_com_varios_itens(tmp_path):
    caminho = tr_sintetico(tmp_path / "tr.docx", [
        ("1", "Papel A4 75g", "Resma", "50"),
        ("2", "Caneta esferográfica azul", "Unidade", "200"),
        ("3", "Grampeador de mesa", "Unidade", "10"),
    ])
    dados = extrair_dados_tr.extrair(caminho, Registro())
    assert [item.numero for item in dados.itens] == ["1", "2", "3"]
    assert [item.quantidade for item in dados.itens] == ["50", "200", "10"]


# --------------------------------------------------------------------------- #
# 5, 6 e 7 — critério de julgamento
# --------------------------------------------------------------------------- #

@pytest.mark.parametrize("criterio", [
    "MENOR PREÇO POR ITEM",
    "MENOR PREÇO POR LOTE",
    "MENOR PREÇO GLOBAL",
])
def test_05_06_07_criterio_de_julgamento_vai_literal_para_o_aviso(tmp_path, criterio):
    manifesto = manifesto_ajustado(
        tmp_path / "manifesto.json", processo__criterio_julgamento=criterio)
    dados = montar_aviso_completo.montar(
        None, manifesto, saida=tmp_path / "saida", converter_pdf=False,
        gerar_pacote=False)
    texto = texto_do(Path(dados["arquivos"][0]).parent / "00_AVISO.docx")
    assert criterio in texto


# --------------------------------------------------------------------------- #
# 8 — modelo de proposta divergente
# --------------------------------------------------------------------------- #

def test_08_modelo_de_proposta_divergente_bloqueia(tmp_path):
    tr_dois = tr_sintetico(tmp_path / "tr2.docx", [
        ("1", "Papel A4 75g", "Resma", "50"),
        ("2", "Caneta esferográfica azul", "Unidade", "200"),
    ])
    dados_dois = extrair_dados_tr.extrair(tr_dois, Registro())
    proposta, _ = gerar_modelo_proposta.gerar(
        MINUTA_PROPOSTA, tmp_path / "proposta.docx", {}, dados_dois, Registro())

    tr_tres = tr_sintetico(tmp_path / "tr3.docx", [
        ("1", "Papel A4 75g", "Resma", "50"),
        ("2", "Caneta esferográfica azul", "Unidade", "200"),
        ("3", "Grampeador de mesa", "Unidade", "10"),
    ])
    registro = Registro()
    gerar_modelo_proposta.conferir_contra_tr(
        proposta, extrair_dados_tr.extrair(tr_tres, Registro()), registro)
    assert registro.tem_bloqueio()
    assert "item(ns)" in registro.bloqueantes[0].mensagem


def test_08b_modelo_de_proposta_nao_traz_preco(montagem_sem_contrato):
    _, saida = montagem_sem_contrato
    proposta = saida / "componentes" / "03_MODELO_DE_PROPOSTA_COMERCIAL.docx"
    documento = Document(str(proposta))
    quadro = [t for t in documento.tables
              if "VALOR UNITÁRIO" in " ".join(c.text for c in t.rows[0].cells)][0]
    for linha in quadro.rows[1:]:
        assert linha.cells[4].text.strip() == ""   # marca
        assert linha.cells[5].text.strip() == ""   # valor unitário
        assert linha.cells[6].text.strip() == ""   # valor total


# --------------------------------------------------------------------------- #
# 9 e 10 — TR pendente e de outro processo
# --------------------------------------------------------------------------- #

def test_09_tr_com_campo_pendente_bloqueia(tmp_path):
    caminho = tr_sintetico(tmp_path / "tr.docx",
                           [("1", "Item qualquer", "Unidade", "1")], pendente=True)
    registro = Registro()
    dados = extrair_dados_tr.extrair(caminho, registro)
    assert dados.campos_pendentes
    assert registro.tem_bloqueio()


def test_10_tr_de_processo_diferente_bloqueia(tmp_path):
    caminho = tr_sintetico(tmp_path / "tr.docx",
                           [("1", "Item qualquer", "Unidade", "1")],
                           processo="099/2025")
    registro = Registro()
    dados = extrair_dados_tr.extrair(caminho, Registro())
    manifesto = json.loads(MANIFESTO.read_text(encoding="utf-8"))
    validar_aviso_completo.validar_identificacao(manifesto, dados, registro)
    assert any("processo diferente" in o.mensagem for o in registro.bloqueantes)


def test_10b_minuta_de_contrato_fora_da_biblioteca_e_recusada(tmp_path):
    falsa = tmp_path / "CONTRATO_MINUTA_MAE.docx"
    shutil.copyfile(MINUTA_CONTRATO, falsa)
    registro = Registro()
    localizar_componentes.decidir_instrumento(
        {"instrumento_contratual": {"tipo": "contrato",
                                    "incluir_minuta_no_aviso": True}},
        registro, minuta_usuario=falsa)
    assert registro.tem_bloqueio()


def test_10c_minuta_de_contrato_nao_cadastrada(tmp_path):
    inexistente = tmp_path / "CONTRATO_DE_OUTRO_ORGAO_MINUTA_MAE.docx"
    inexistente.write_bytes(b"nao e docx")
    registro = Registro()
    localizar_componentes.decidir_instrumento(
        {"instrumento_contratual": {"tipo": "contrato",
                                    "incluir_minuta_no_aviso": True}},
        registro, minuta_usuario=inexistente)
    assert any(localizar_componentes.MSG_SEM_MINUTA_CONTRATO in o.mensagem
               for o in registro.bloqueantes)


def test_10d_instrumento_indefinido_bloqueia():
    registro = Registro()
    instrumento = localizar_componentes.decidir_instrumento({}, registro)
    assert not instrumento.definido
    assert any("PENDÊNCIA: definir se a contratação" in o.mensagem
               for o in registro.bloqueantes)


# --------------------------------------------------------------------------- #
# 11 e 12 — numeração da declaração
# --------------------------------------------------------------------------- #

def test_11_declaracao_e_anexo_iv_sem_contrato(montagem_sem_contrato):
    dados, _ = montagem_sem_contrato
    declaracao = [a for a in dados["anexos"] if a["componente"] == "declaracao"][0]
    assert declaracao["romano"] == "IV"
    assert declaracao["rotulo"] == "ANEXO IV — DECLARAÇÃO CONJUNTA"


def test_12_declaracao_e_anexo_v_com_contrato():
    registro = Registro()
    manifesto = json.loads(MANIFESTO_CONTRATO.read_text(encoding="utf-8"))
    componentes = localizar_componentes.localizar(
        manifesto, registro, tr_explicito=TR_FIXTURE)
    plano = numerar_anexos.planejar(componentes, registro)
    assert plano.por_chave("declaracao").romano == "V"


# --------------------------------------------------------------------------- #
# 13 e 14 — relação de anexos e referências internas
# --------------------------------------------------------------------------- #

def _plano_de_quatro():
    registro = Registro()
    manifesto = json.loads(MANIFESTO.read_text(encoding="utf-8"))
    componentes = localizar_componentes.localizar(
        manifesto, registro, tr_explicito=TR_FIXTURE)
    return numerar_anexos.planejar(componentes, registro)


def test_13_relacao_de_anexos_e_atualizada(tmp_path):
    documento = Document()
    documento.add_paragraph("1.3 Compõem este Aviso os seguintes documentos:")
    documento.add_paragraph("ANEXO I – Documentação exigida para Habilitação;")
    documento.add_paragraph("ANEXO II – Termo de Referência;")
    documento.add_paragraph("ANEXO III – Modelo de Proposta de Preços;")
    documento.add_paragraph("ANEXO IV – Declaração Conjunta;")

    registro = Registro()
    correcoes = numerar_anexos.atualizar_relacao(documento, _plano_de_quatro(), registro)
    textos = [p.text for p in documento.paragraphs]
    assert "ANEXO II — TERMO DE REFERÊNCIA;" in textos
    assert "ANEXO IV — DECLARAÇÃO CONJUNTA;" in textos
    assert correcoes


def test_13b_ausencia_de_relacao_vira_pendencia_e_nao_invencao():
    documento = Document(str(MINUTA_AVISO))
    registro = Registro()
    correcoes = numerar_anexos.atualizar_relacao(documento, _plano_de_quatro(), registro)
    assert correcoes == []
    assert registro.pendencias
    assert not registro.tem_bloqueio()


def test_14_referencia_a_anexo_inexistente_bloqueia():
    documento = Document()
    documento.add_paragraph("Observar o disposto no ANEXO VII deste aviso.")
    registro = Registro()
    numerar_anexos.verificar_referencias(
        documento.paragraphs, _plano_de_quatro(), registro)
    assert any("ANEXO VII" in o.mensagem for o in registro.bloqueantes)


def test_14b_numeracao_sem_lacuna_nem_duplicidade(montagem_sem_contrato):
    dados, _ = montagem_sem_contrato
    romanos = [a["romano"] for a in dados["anexos"]]
    assert romanos == sorted(set(romanos), key=romanos.index)
    assert romanos == ["I", "II", "III", "IV"]


# --------------------------------------------------------------------------- #
# 15, 16, 17 e 18 — timbre, tabelas, imagens e paisagem
# --------------------------------------------------------------------------- #

def test_15_cabecalho_e_rodape_preservados(montagem_sem_contrato):
    _, saida = montagem_sem_contrato
    final = saida / "saida" / "AVISO_DISPENSA_COMPLETO.docx"
    intacto, divergencias = unir_docx.timbre_intacto(MINUTA_AVISO, final)
    assert intacto, divergencias


def test_15b_documento_final_tem_um_unico_timbre(montagem_sem_contrato):
    _, saida = montagem_sem_contrato
    final = saida / "saida" / "AVISO_DISPENSA_COMPLETO.docx"
    with zipfile.ZipFile(final) as pacote:
        cabecalhos = [n for n in pacote.namelist()
                      if n.startswith("word/header")and n.endswith(".xml")]
    assert len(cabecalhos) == 1


def test_16_tabela_extensa_sobrevive_a_montagem(tmp_path):
    itens = [(str(i), f"Item de teste número {i}", "Unidade", str(i * 2))
             for i in range(1, 31)]
    caminho = tr_sintetico(tmp_path / "tr.docx", itens)
    dados = extrair_dados_tr.extrair(caminho, Registro())
    proposta, _ = gerar_modelo_proposta.gerar(
        MINUTA_PROPOSTA, tmp_path / "proposta.docx", {}, dados, Registro())

    documento = Document(str(proposta))
    quadro = [t for t in documento.tables
              if "QNTD" in " ".join(c.text for c in t.rows[0].cells).upper()][0]
    assert len(quadro.rows) == 31
    assert quadro.rows[30].cells[1].text.strip() == "Item de teste número 30"


def test_17_imagens_preservadas(montagem_sem_contrato):
    _, saida = montagem_sem_contrato
    final = saida / "saida" / "AVISO_DISPENSA_COMPLETO.docx"
    with zipfile.ZipFile(MINUTA_AVISO) as origem, zipfile.ZipFile(final) as destino:
        do_aviso = {n for n in origem.namelist() if n.startswith("word/media/")}
        do_final = {n for n in destino.namelist() if n.startswith("word/media/")}
    assert len(do_final) >= len(do_aviso) - 1  # órfãs somem ao salvar; timbre fica
    assert do_final


def test_18_secao_em_paisagem_preserva_orientacao(tmp_path):
    paisagem = Document()
    secao = paisagem.sections[0]
    secao.orientation = WD_ORIENT.LANDSCAPE
    secao.page_width, secao.page_height = Cm(29.7), Cm(21)
    paisagem.add_paragraph("Quadro em paisagem — conteúdo de teste.")
    caminho = tmp_path / "paisagem.docx"
    paisagem.save(str(caminho))

    destino = tmp_path / "unido.docx"
    unir_docx.unir(MINUTA_AVISO, [caminho], destino, Registro())
    documento = Document(str(destino))
    assert any(s.orientation == WD_ORIENT.LANDSCAPE for s in documento.sections)
    assert "Quadro em paisagem" in texto_do(destino)


# --------------------------------------------------------------------------- #
# 19, 20 e 21 — runs, idempotência e preservação
# --------------------------------------------------------------------------- #

def test_19_marcador_dividido_em_runs(tmp_path):
    documento = Document()
    paragrafo = documento.add_paragraph()
    for pedaco in ("Objeto: {{OB", "JE", "TO}} — fim"):
        paragrafo.add_run(pedaco)
    caminho = tmp_path / "dividido.docx"
    documento.save(str(caminho))

    montar_aviso_completo.preencher_documento(
        caminho, {"OBJETO": "Aquisição de material de limpeza"}, Registro(), "teste")
    texto = texto_do(caminho)
    assert "Aquisição de material de limpeza" in texto
    assert "{{" not in texto


def test_20_montagem_executada_duas_vezes(tmp_path):
    primeira = montar_aviso_completo.montar(
        None, MANIFESTO, saida=tmp_path / "saida", converter_pdf=False,
        gerar_pacote=False)
    segunda = montar_aviso_completo.montar(
        None, MANIFESTO, saida=tmp_path / "saida", converter_pdf=False,
        gerar_pacote=False)

    assert primeira["status"] == segunda["status"]
    assert primeira["anexos"] == segunda["anexos"]
    assert primeira["contagem"] == segunda["contagem"]


def test_21_conteudo_integralmente_preservado(montagem_sem_contrato):
    dados, saida = montagem_sem_contrato
    assert dados["validacao_conteudo"]["Linhas perdidas"] == 0
    assert dados["validacao_conteudo"]["Linhas conferidas"] > 100

    final = texto_do(saida / "saida" / "AVISO_DISPENSA_COMPLETO.docx")
    for trecho in (
        "prova de regularidade com o Fundo de Garantia do Tempo de Serviço",
        "DECLARA a inexistência de fato impeditivo",
        "forno micro-ondas",
    ):
        assert trecho in final


def test_21b_campos_do_aviso_preenchidos(montagem_sem_contrato):
    dados, saida = montagem_sem_contrato
    texto = texto_do(saida / "saida" / "AVISO_DISPENSA_COMPLETO.docx")
    assert "011/2026" in texto
    assert "27/07/2026, às 08:00" in texto
    assert "art. 75, inciso II" in texto          # sem "art. art."
    assert "art. art." not in texto
    assert not dados["campos_pendentes"]


# --------------------------------------------------------------------------- #
# 22 a 25 — PDF, ZIP e abertura nos editores
# --------------------------------------------------------------------------- #

@pytest.mark.slow
def test_22_conversao_para_pdf(tmp_path):
    conversor = gerar_pacote_publicacao.detectar_conversor()
    if not conversor.disponivel():
        pytest.skip("nenhum conversor de PDF disponível no ambiente")
    destino = tmp_path / "aviso.pdf"
    gerado = gerar_pacote_publicacao.converter_pdf(
        MINUTA_AVISO, destino, conversor, Registro())
    assert gerado is not None and gerado.exists()
    assert gerado.read_bytes().startswith(b"%PDF")


def test_23_pacote_zip(montagem_sem_contrato):
    _, saida = montagem_sem_contrato
    caminho = saida / "saida" / "PACOTE_PUBLICACAO.zip"
    assert caminho.exists()
    with zipfile.ZipFile(caminho) as pacote:
        nomes = pacote.namelist()
    assert "MANIFESTO_ARQUIVOS.json" in nomes
    assert "AVISO_DISPENSA_COMPLETO.docx" in nomes
    assert "01_AVISO_DE_CONTRATACAO_DIRETA.docx" in nomes
    assert "05_ANEXO_IV_DECLARACAO_CONJUNTA.docx" in nomes
    assert not any("ANEXO_V_" in n for n in nomes)   # sem contrato, não há Anexo V


def test_23b_manifesto_de_arquivos_confere_com_os_anexos(montagem_sem_contrato):
    dados, saida = montagem_sem_contrato
    manifesto = json.loads(
        (saida / "saida" / "MANIFESTO_ARQUIVOS.json").read_text(encoding="utf-8"))
    assert len(manifesto["anexos"]) == len(dados["anexos"])
    assert [a["rotulo"] for a in manifesto["anexos"]] == \
           [a["rotulo"] for a in dados["anexos"]]


@pytest.mark.slow
def test_24_abertura_no_microsoft_word(montagem_sem_contrato):
    import subprocess
    import sys as _sys

    if _sys.platform != "win32":
        pytest.skip("Microsoft Word só é verificável no Windows")
    conversor = gerar_pacote_publicacao.detectar_conversor()
    if conversor.nome != "word":
        pytest.skip("Microsoft Word não disponível neste ambiente")

    _, saida = montagem_sem_contrato
    final = (saida / "saida" / "AVISO_DISPENSA_COMPLETO.docx").resolve()
    script = (
        "$app = New-Object -ComObject Word.Application; $app.Visible = $false; "
        f"$doc = $app.Documents.Open('{final}', $false, $true); "
        "Write-Output $doc.Paragraphs.Count; $doc.Close(0); $app.Quit()"
    )
    processo = subprocess.run(
        ["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
        capture_output=True, text=True, timeout=300)
    assert processo.returncode == 0, processo.stderr
    assert int(processo.stdout.strip().splitlines()[-1]) > 100


@pytest.mark.slow
def test_25_abertura_no_libreoffice(montagem_sem_contrato):
    import subprocess

    conversor = gerar_pacote_publicacao.detectar_conversor()
    if conversor.nome != "libreoffice":
        pytest.skip("LibreOffice não disponível neste ambiente")

    _, saida = montagem_sem_contrato
    final = saida / "saida" / "AVISO_DISPENSA_COMPLETO.docx"
    processo = subprocess.run(
        [conversor.caminho, "--headless", "--convert-to", "pdf",
         "--outdir", str(saida / "libreoffice"), str(final)],
        capture_output=True, timeout=300)
    assert processo.returncode == 0
    assert (saida / "libreoffice" / "AVISO_DISPENSA_COMPLETO.pdf").exists()


# --------------------------------------------------------------------------- #
# Regras que não podem regredir
# --------------------------------------------------------------------------- #

def test_regra_nao_grava_em_minutas(tmp_path):
    """A montagem lê a biblioteca oficial; nunca escreve nela."""
    antes = {caminho: caminho.stat().st_mtime
             for caminho in (RAIZ / "05_minutas").rglob("*.docx")}
    montar_aviso_completo.montar(
        None, MANIFESTO, saida=tmp_path / "saida", converter_pdf=False,
        gerar_pacote=False)
    depois = {caminho: caminho.stat().st_mtime
              for caminho in (RAIZ / "05_minutas").rglob("*.docx")}
    assert antes == depois


def test_regra_prazo_minimo_de_tres_dias_uteis(tmp_path):
    manifesto = manifesto_ajustado(
        tmp_path / "manifesto.json",
        recebimento_propostas__data_inicio="2026-07-27",
        recebimento_propostas__data_fim="2026-07-28")
    registro = Registro()
    validar_aviso_completo.validar_datas(
        json.loads(manifesto.read_text(encoding="utf-8")), registro)
    assert any("dia(s) útil(eis)" in o.mensagem for o in registro.bloqueantes)


def test_regra_data_final_anterior_a_inicial(tmp_path):
    manifesto = manifesto_ajustado(
        tmp_path / "manifesto.json",
        recebimento_propostas__data_inicio="2026-07-29",
        recebimento_propostas__data_fim="2026-07-27")
    registro = Registro()
    validar_aviso_completo.validar_datas(
        json.loads(manifesto.read_text(encoding="utf-8")), registro)
    assert registro.tem_bloqueio()


def test_regra_exigencia_negada_no_tr_nao_vira_divergencia():
    """"Não será exigida qualificação econômico-financeira" não é exigência."""
    dados = extrair_dados_tr.extrair(TR_FIXTURE, Registro())
    assert dados.sinais["qualificacao_economica"] is False
    assert dados.sinais["qualificacao_tecnica"] is False

    registro = Registro()
    validar_aviso_completo.validar_habilitacao(
        texto_do(MINUTA_AVISO), dados, registro)
    assert not registro.pendencias


def test_regra_status_apto_para_publicacao_nao_e_automatico():
    import relatorio_aviso_completo
    assert relatorio_aviso_completo.status_permitido(
        "APTO PARA PUBLICAÇÃO") == "APTO PARA CONFERÊNCIA"


def test_relatorio_tem_todas_as_secoes_do_escopo(montagem_sem_contrato):
    dados, saida = montagem_sem_contrato
    markdown = (saida / "relatorios" / "VALIDACAO_AVISO_COMPLETO.md").read_text(
        encoding="utf-8")
    for titulo in ("Identificação do processo", "Componentes localizados",
                   "Minutas oficiais utilizadas", "Termo de Referência utilizado",
                   "Instrumento contratual definido", "Ordem dos anexos",
                   "Dados extraídos do TR", "Validação da habilitação",
                   "Validação do modelo de proposta",
                   "Validação da minuta de contrato",
                   "Validação da declaração conjunta", "Divergências encontradas",
                   "Campos pendentes", "Correções realizadas",
                   "Validação de conteúdo", "Validação de formatação",
                   "Arquivos produzidos", "Resultado"):
        assert f"## {titulo}" in markdown
    assert "Conferência humana final obrigatória" in markdown
    assert (saida / "relatorios" / "VALIDACAO_AVISO_COMPLETO.json").exists()


def test_sem_folha_em_branco_entre_o_aviso_e_o_anexo_i(tmp_path):
    """
    O Anexo I vinha precedido de uma quebra de página herdada da minuta. Somada
    à quebra de seção que fecha o aviso, ela produzia uma folha em branco entre
    a assinatura do Presidente e o Anexo I.
    """
    from util_ooxml import quebras_de_pagina

    aviso = tmp_path / "aviso.docx"
    habilitacao = tmp_path / "habilitacao.docx"
    _, gerado = unir_docx.dividir_aviso(
        localizar_componentes.MINUTA_AVISO, aviso, habilitacao, Registro())

    assert gerado is not None
    primeiro = Document(str(gerado)).paragraphs[0]
    assert quebras_de_pagina(primeiro) == 0


def test_ultimo_anexo_nao_abre_secao_vazia(tmp_path):
    """Quebra de seção no último componente deixaria uma folha em branco no fim."""
    from util_ooxml import qn

    partes = []
    for indice in range(2):
        documento = Document()
        documento.add_paragraph(f"Conteúdo do componente {indice}.")
        caminho = tmp_path / f"parte{indice}.docx"
        documento.save(str(caminho))
        partes.append(caminho)

    destino = tmp_path / "unido.docx"
    unir_docx.unir(localizar_componentes.MINUTA_AVISO, partes, destino, Registro())

    documento = Document(str(destino))
    corpo = documento.element.body
    # Uma seção por componente: o mestre e a primeira parte fecham a sua; a
    # última é fechada pelo sectPr de corpo. Nenhuma seção a mais.
    assert len(list(corpo.iter(qn("w:sectPr")))) == len(partes) + 1


def test_todas_as_secoes_usam_as_margens_do_aviso(tmp_path):
    """
    Componente com margens de documento em branco jogava o texto por cima do
    rodapé timbrado. O tamanho da página continua sendo o do componente; as
    margens passam a ser as do aviso, que reservam espaço para o timbre.
    """
    from util_ooxml import qn

    simples = Document()
    simples.add_paragraph("Componente com margens padrão do Word.")
    caminho = tmp_path / "simples.docx"
    simples.save(str(caminho))

    destino = tmp_path / "unido.docx"
    unir_docx.unir(localizar_componentes.MINUTA_AVISO, [caminho], destino,
                   Registro())

    def margens(elemento):
        pgMar = elemento.find(qn("w:pgMar"))
        return None if pgMar is None else {
            k.split("}")[-1]: v for k, v in pgMar.attrib.items()}

    corpo = Document(str(destino)).element.body
    encontradas = [margens(s) for s in corpo.iter(qn("w:sectPr"))]
    assert encontradas and all(m == encontradas[0] for m in encontradas)


def test_quadro_de_itens_cabe_na_largura_util(tmp_path):
    """
    A especificação do TR vem inteira na célula de descrição. Com as colunas
    repartidas por igual, um único item ocupava três páginas.
    """
    caminho = tr_sintetico(tmp_path / "tr.docx", [
        ("1", "Especificação longa. " * 60, "Unidade", "1"),
    ])
    dados = extrair_dados_tr.extrair(caminho, Registro())
    destino, _ = gerar_modelo_proposta.gerar(
        MINUTA_PROPOSTA, tmp_path / "proposta.docx", {}, dados, Registro())

    documento = Document(str(destino))
    quadro = [t for t in documento.tables
              if "DESCRIÇÃO" in " ".join(c.text for c in t.rows[0].cells).upper()][0]
    larguras = [c.width for c in quadro.columns]
    assert all(l is not None for l in larguras)
    # A descrição é a coluna larga; nenhuma outra chega perto dela.
    assert larguras[1] == max(larguras)
    assert larguras[1] > sum(larguras) * 0.35


@pytest.mark.parametrize("minuta", [
    localizar_componentes.MINUTA_PROPOSTA,
    localizar_componentes.MINUTA_DECLARACAO,
])
def test_minutas_de_terceiros_exibem_o_timbre_oficial(minuta):
    """
    Até a v1.1 estas minutas não exibiam timbre: o cabeçalho e o rodapé estavam
    no pacote, mas o `sectPr` não os referenciava — existiam e nunca apareciam.
    Passaram a usar o timbre oficial do aviso.
    """
    documento = Document(str(minuta))
    secao = documento.sections[0]
    assert not secao.header.is_linked_to_previous, "cabeçalho não referenciado"
    assert not secao.footer.is_linked_to_previous, "rodapé não referenciado"

    intacto, divergencias = unir_docx.timbre_intacto(MINUTA_AVISO, minuta)
    assert intacto, divergencias


def test_montagem_nao_acusa_mais_timbre_divergente(montagem_sem_contrato):
    """Com as minutas corrigidas, não há mais alerta de timbre na montagem."""
    dados, _ = montagem_sem_contrato
    alertas = [o["mensagem"] for o in dados["ocorrencias"]
               if o["severidade"] == "ALERTA"]
    assert not any("cabeçalho/rodapé diferente" in m for m in alertas), alertas


@pytest.mark.parametrize("minuta", [
    localizar_componentes.MINUTA_PROPOSTA,
    localizar_componentes.MINUTA_DECLARACAO,
])
def test_minutas_do_fornecedor_so_guardam_marcadores_da_administracao(minuta):
    """
    São formulários preenchidos por terceiros. O que cabe ao fornecedor tem de
    aparecer como campo — régua ou célula em branco —, nunca como `{{CAMPO}}` ou
    `[PREENCHER]`, que não dizem a ele onde escrever.
    """
    import re

    from util_ooxml import CAMPO_RE

    texto = texto_do(minuta)
    restantes = set(CAMPO_RE.findall(texto))
    permitidos = {"NUMERO_PROCESSO", "NUMERO_AVISO", "OBJETO",
                  "ITEM", "DESCRICAO_ITEM", "UNIDADE", "QUANTIDADE"}
    assert restantes <= permitidos, restantes - permitidos
    assert "[PREENCHER" not in texto.upper()
    assert re.search(r"_{4,}", texto), "nenhum campo de preenchimento desenhado"


def test_quadro_de_itens_da_proposta_e_emoldurado():
    """Célula em branco sem borda não se lê como campo."""
    from util_ooxml import qn

    documento = Document(str(localizar_componentes.MINUTA_PROPOSTA))
    assert documento.tables
    for tabela in documento.tables:
        assert tabela._tbl.tblPr.find(qn("w:tblBorders")) is not None
