#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
montar_aviso_completo.py — comando do Aviso de Dispensa Completo.

Executa o fluxo do item 6 do escopo, do manifesto ao pacote de publicação:

    manifesto → componentes → TR → proposta → contrato → declaração →
    numeração → documento único → padronização → validação → relatório → pacote

Exemplos:
    python scripts/aviso_completo/montar_aviso_completo.py \
      --processo 08_processos_em_andamento/PA_031_2026/

    python scripts/aviso_completo/montar_aviso_completo.py \
      --aviso 05_minutas/AVISO/AVISO_CONTRATACAO_DIRETA_MINUTA_MAE.docx \
      --tr caminho/TR_FINAL.docx --contrato 05_minutas/CONTRATO/CONTRATO_MINUTA_MAE.docx \
      --incluir-contrato --manifesto manifesto_aviso_completo.json \
      --saida 08_processos_em_andamento/PA_031_2026/07_AVISO_COMPLETO/saida/

O arquivo de entrada nunca é sobrescrito e nada é gravado em `05_minutas/`.
"""
from __future__ import annotations

import argparse
import shutil
import sys
import tempfile
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

if __package__ in (None, ""):
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "docx_cmi"))

import extrair_dados_tr
import gerar_modelo_proposta
import gerar_pacote_publicacao
import localizar_componentes
import numerar_anexos
import relatorio_aviso_completo
import unir_docx
import validar_aviso_completo
from numerar_anexos import PlanoAnexos
from ocorrencias import (
    AVISO_RASCUNHO,
    STATUS_RASCUNHO,
    Registro,
    decidir_status,
)
from util_ooxml import (
    CAMPO_RE,
    exigir_python_docx,
    iter_paragrafos_corpo,
    substituir_marcadores_paragrafo,
    texto_paragrafo,
)
import formatar_docx

RAIZ = Path(__file__).resolve().parents[2]
ETAPA = "montagem"

PASTA_SAIDA_PADRAO = "07_AVISO_COMPLETO"

MESES = ("janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho",
         "agosto", "setembro", "outubro", "novembro", "dezembro")


# --------------------------------------------------------------------------- #
# Campos do aviso
# --------------------------------------------------------------------------- #

def _data_extenso(valor: Optional[str]) -> str:
    referencia = date.today()
    if valor:
        try:
            referencia = datetime.strptime(str(valor).strip(), "%Y-%m-%d").date()
        except ValueError:
            return str(valor).strip()
    return f"{referencia.day} de {MESES[referencia.month - 1]} de {referencia.year}"


def _data_hora(data: Optional[str], hora: Optional[str]) -> str:
    if not data:
        return ""
    try:
        formatada = datetime.strptime(str(data).strip(), "%Y-%m-%d").strftime("%d/%m/%Y")
    except ValueError:
        formatada = str(data).strip()
    return f"{formatada}, às {hora}" if hora else formatada


def campos_do_aviso(manifesto: dict[str, Any], registro: Registro) -> dict[str, str]:
    """
    Monta os valores dos campos `{{...}}` da minuta do aviso.

    O fundamento legal recebe um ajuste explícito: a minuta já escreve "art."
    antes do marcador, e o manifesto costuma trazer "art. 75, inciso II...".
    Sem o ajuste sairia "art. art. 75" no documento publicado.
    """
    processo = manifesto.get("processo", {})
    recebimento = manifesto.get("recebimento_propostas", {})
    assinatura = manifesto.get("assinatura", {})

    fundamento = str(processo.get("fundamento_legal") or "").strip()
    sem_art = fundamento
    for prefixo in ("art. ", "art.", "artigo "):
        if sem_art.lower().startswith(prefixo):
            sem_art = sem_art[len(prefixo):].strip()
            break
    if sem_art != fundamento:
        registro.informacao(
            ETAPA,
            f"Fundamento legal ajustado ao texto da minuta: '{fundamento}' → "
            f"'{sem_art}' (a minuta já traz 'art.' antes do campo).",
        )

    return {
        "NUMERO_AVISO": str(processo.get("dispensa") or ""),
        "NUMERO_PROCESSO": str(processo.get("numero") or ""),
        "OBJETO": str(processo.get("objeto") or ""),
        "CRITERIO_JULGAMENTO": str(processo.get("criterio_julgamento") or ""),
        "FUNDAMENTO_LEGAL": sem_art,
        "DATA_INICIO_PROPOSTAS": _data_hora(recebimento.get("data_inicio"),
                                            recebimento.get("hora_inicio")),
        "DATA_FIM_PROPOSTAS": _data_hora(recebimento.get("data_fim"),
                                         recebimento.get("hora_fim")),
        "DATA": _data_extenso(assinatura.get("data")),
        "NOME_PRESIDENTE": str(assinatura.get("nome") or ""),
        "CARGO": str(assinatura.get("cargo") or ""),
    }


def preencher_documento(caminho: Path, campos: dict[str, str],
                        registro: Registro, rotulo: str) -> list[str]:
    """Preenche `{{CAMPO}}` no documento, no lugar, preservando a formatação."""
    exigir_python_docx()
    from docx import Document

    documento = Document(str(caminho))
    preenchidos: list[str] = []

    def resolver(nome: str) -> Optional[str]:
        valor = campos.get(nome)
        if valor is None or valor == "":
            return None
        preenchidos.append(nome)
        return valor

    for paragrafo in iter_paragrafos_corpo(documento):
        substituir_marcadores_paragrafo(paragrafo, resolver, CAMPO_RE)
    documento.save(str(caminho))

    nomes = sorted(set(preenchidos))
    registro.informacao(
        ETAPA, f"{rotulo}: campos preenchidos — {', '.join(nomes) or 'nenhum'}.",
        origem=caminho.name)
    return nomes


# --------------------------------------------------------------------------- #
# Componentes
# --------------------------------------------------------------------------- #

def _texto_do_documento(caminho: Path) -> str:
    exigir_python_docx()
    from docx import Document
    return "\n".join(
        texto_paragrafo(p).strip()
        for p in iter_paragrafos_corpo(Document(str(caminho)))
    )


def montar_componentes(componentes: localizar_componentes.Componentes,
                       plano: PlanoAnexos, dados_tr, campos: dict[str, str],
                       pasta_componentes: Path,
                       registro: Registro) -> tuple[dict[str, Path], list[str]]:
    """
    Produz um DOCX por componente, já preenchido, rotulado e paginado.

    Devolve também as linhas com lacunas reservadas ao proponente, para que a
    validação final não as confunda com campo esquecido.
    """
    pasta_componentes.mkdir(parents=True, exist_ok=True)
    produzidos: dict[str, Path] = {}
    lacunas: list[str] = []

    aviso = componentes.obter("aviso")
    destino_aviso = pasta_componentes / "00_AVISO.docx"
    anexo_habilitacao = plano.por_chave("habilitacao")
    destino_habilitacao = pasta_componentes / (
        anexo_habilitacao.nome_componente if anexo_habilitacao
        else "01_HABILITACAO.docx")

    caminho_aviso, caminho_habilitacao = unir_docx.dividir_aviso(
        aviso.caminho, destino_aviso, destino_habilitacao, registro)
    produzidos["aviso"] = caminho_aviso
    if caminho_habilitacao is not None:
        produzidos["habilitacao"] = caminho_habilitacao

    preencher_documento(caminho_aviso, campos, registro, "Aviso")
    if caminho_habilitacao is not None:
        preencher_documento(caminho_habilitacao, campos, registro, "Anexo I")

    dados_anexos = {
        "NUMERO_PROCESSO": campos.get("NUMERO_PROCESSO", ""),
        "NUMERO_AVISO": campos.get("NUMERO_AVISO", ""),
        "OBJETO": campos.get("OBJETO", ""),
        "CRITERIO_JULGAMENTO": campos.get("CRITERIO_JULGAMENTO", ""),
        "FUNDAMENTO_LEGAL": campos.get("FUNDAMENTO_LEGAL", ""),
    }

    for anexo in plano.anexos:
        if anexo.chave == "habilitacao":
            continue
        destino = pasta_componentes / anexo.nome_componente
        origem = anexo.componente.caminho

        if anexo.chave == "tr":
            shutil.copyfile(origem, destino)
        elif anexo.chave == "proposta":
            _, novas = gerar_modelo_proposta.gerar(
                origem, destino, dados_anexos, dados_tr, registro)
            lacunas += novas
        elif anexo.chave == "contrato":
            _, novas = gerar_modelo_proposta.gerar_minuta_contrato(
                origem, destino, dados_anexos, registro)
            lacunas += novas
        elif anexo.chave == "declaracao":
            _, novas = gerar_modelo_proposta.gerar_declaracao(
                origem, destino, dados_anexos, registro)
            lacunas += novas
        else:  # pragma: no cover - toda chave conhecida está tratada acima
            shutil.copyfile(origem, destino)

        unir_docx.inserir_titulo_anexo(destino, destino, anexo.rotulo, registro)
        produzidos[anexo.chave] = destino

    return produzidos, lacunas


# --------------------------------------------------------------------------- #
# Fluxo principal
# --------------------------------------------------------------------------- #

def montar(pasta_processo: Optional[Path], caminho_manifesto: Optional[Path],
           saida: Optional[Path] = None, aviso: Optional[Path] = None,
           tr: Optional[Path] = None, contrato: Optional[Path] = None,
           decisao_contrato: Optional[bool] = None,
           autorizar_tr_rascunho: bool = False,
           permitir_rascunho: bool = False,
           padronizar: bool = True, gerar_pacote: bool = True,
           converter_pdf: bool = True) -> dict[str, Any]:
    """Executa a montagem completa e devolve os dados do relatório."""
    registro = Registro()

    if caminho_manifesto is None and pasta_processo is not None:
        caminho_manifesto = localizar_componentes.descobrir_manifesto(pasta_processo)
    if caminho_manifesto is None:
        registro.bloqueante(
            ETAPA,
            "Manifesto do aviso completo não encontrado. Informe --manifesto ou "
            f"crie {localizar_componentes.MANIFESTO_PADRAO} na pasta do processo.")
        return _dados_relatorio(registro, {}, None, None, None, {}, {},
                                permitir_rascunho)

    manifesto = localizar_componentes.carregar_manifesto(caminho_manifesto)

    # Pré-leitura do TR só para decidir o instrumento contratual. O registro
    # descartado evita que a mesma falha seja reportada duas vezes.
    texto_tr = ""
    prescan = Registro()
    caminho_tr = tr or localizar_componentes.localizar_tr(
        manifesto, pasta_processo, prescan)
    if caminho_tr is not None and Path(caminho_tr).exists():
        try:
            texto_tr = _texto_do_documento(Path(caminho_tr))
        except Exception:  # pragma: no cover - TR ilegível vira bloqueio adiante
            texto_tr = ""

    componentes = localizar_componentes.localizar(
        manifesto, registro,
        pasta_processo=pasta_processo,
        caminho_manifesto=caminho_manifesto,
        tr_explicito=Path(caminho_tr) if caminho_tr else None,
        aviso_explicito=aviso,
        contrato_explicito=contrato,
        decisao_contrato=decisao_contrato,
        autorizar_tr_rascunho=autorizar_tr_rascunho,
        texto_tr=texto_tr,
    )

    componente_tr = componentes.obter("tr")
    dados_tr = None
    if componente_tr is not None and componente_tr.caminho is not None:
        dados_tr = extrair_dados_tr.extrair(componente_tr.caminho, registro)

    validar_aviso_completo.validar_identificacao(manifesto, dados_tr, registro)
    validar_aviso_completo.validar_datas(manifesto, registro)

    plano = numerar_anexos.planejar(componentes, registro)
    campos = campos_do_aviso(manifesto, registro)

    pasta_base = _pasta_de_saida(pasta_processo, saida, caminho_manifesto)
    pasta_componentes = pasta_base / "componentes"
    pasta_saida = pasta_base / "saida"
    pasta_relatorios = pasta_base / "relatorios"

    if registro.tem_bloqueio() and not permitir_rascunho:
        return _dados_relatorio(registro, manifesto, componentes, dados_tr, plano,
                                {}, {}, permitir_rascunho,
                                relatorio_em=pasta_relatorios)

    produzidos, lacunas = montar_componentes(componentes, plano, dados_tr, campos,
                                             pasta_componentes, registro)

    aviso_pronto = produzidos.get("aviso")
    if aviso_pronto is not None:
        exigir_python_docx()
        from docx import Document
        documento = Document(str(aviso_pronto))
        correcoes = numerar_anexos.atualizar_relacao(documento, plano, registro)
        if correcoes:
            documento.save(str(aviso_pronto))
        validar_aviso_completo.validar_habilitacao(
            _texto_do_documento(produzidos.get("habilitacao", aviso_pronto)),
            dados_tr, registro)
    else:
        correcoes = []

    if "proposta" in produzidos:
        gerar_modelo_proposta.conferir_contra_tr(
            produzidos["proposta"], dados_tr, registro)

    ordem = ["aviso"] + [a.chave for a in plano.anexos]
    partes = [produzidos[chave] for chave in ordem if chave in produzidos]
    unir_docx.conferir_timbres(componentes.obter("aviso").caminho, partes, registro)

    pasta_saida.mkdir(parents=True, exist_ok=True)
    documento_final = pasta_saida / "AVISO_DISPENSA_COMPLETO.docx"
    formatacao: dict[str, Any] = {}

    with tempfile.TemporaryDirectory() as temporario:
        bruto = Path(temporario) / "unido.docx"
        unir_docx.unir(partes[0], partes[1:], bruto, registro)

        if padronizar:
            formatacao = _padronizar(bruto, documento_final, registro)
        if not padronizar or not formatacao.get("gravado"):
            shutil.copyfile(bruto, documento_final)

    validacao_final = validar_aviso_completo.validar_documento_final(
        documento_final, plano, componentes.obter("aviso").caminho, registro,
        lacunas_autorizadas=lacunas)
    preservacao = validar_aviso_completo.validar_preservacao(
        partes, documento_final, registro,
        acrescimos_autorizados=[a.rotulo for a in plano.anexos] + [AVISO_RASCUNHO])
    pendencias = validar_aviso_completo.levantar_pendencias(documento_final)

    status = decidir_status(registro, rascunho_solicitado=permitir_rascunho)
    if status == STATUS_RASCUNHO:
        unir_docx.marcar_rascunho(documento_final, AVISO_RASCUNHO, registro)

    relatorio_md = pasta_relatorios / "VALIDACAO_AVISO_COMPLETO.md"
    pacote = None
    if gerar_pacote and not registro.tem_bloqueio():
        pacote = gerar_pacote_publicacao.gerar(
            documento_final, produzidos, plano, pasta_saida, relatorio_md,
            registro, converter=converter_pdf)
    elif gerar_pacote:
        registro.informacao(
            ETAPA,
            "Pacote de publicação não gerado: há erro bloqueante em aberto.")

    dados = _dados_relatorio(
        registro, manifesto, componentes, dados_tr, plano, produzidos,
        {
            "documento_final": str(documento_final),
            "formatacao": formatacao,
            "validacao_final": validacao_final,
            "preservacao": preservacao,
            "pendencias": pendencias,
            "pacote": pacote.como_dicionario() if pacote else None,
            "correcoes": correcoes,
        },
        permitir_rascunho, relatorio_em=pasta_relatorios,
    )
    if pacote is not None and dados.get("relatorio_markdown"):
        gerar_pacote_publicacao.anexar_relatorio(
            pacote.zip, Path(dados["relatorio_markdown"]), registro)
    return dados


def _padronizar(entrada: Path, destino: Path, registro: Registro) -> dict[str, Any]:
    """
    Aplica o Módulo de Padronização Documental ao documento montado.

    A renumeração fica desligada de propósito: o documento único reúne artigos,
    cláusulas e itens de TR, e renumerar qualquer um deles alteraria conteúdo
    jurídico. Aqui só entra acabamento visual, com preservação de conteúdo
    obrigatória — se a padronização bloquear, vale o documento sem ela.
    """
    try:
        diagnostico = formatar_docx.formatar(
            entrada=entrada, saida=destino, perfil_id="aviso",
            corrigir_numeracao=False, preservar_conteudo=True,
            verificar_idempotencia=True,
        )
    except Exception as erro:  # pragma: no cover - depende do documento
        registro.alerta(
            ETAPA,
            f"A padronização documental não pôde ser aplicada ({erro}). O "
            "documento único foi gerado sem o acabamento automático.",
        )
        return {"gravado": False, "erro": str(erro)}

    gravado = bool(diagnostico.get("arquivo_final"))
    if not gravado:
        registro.alerta(
            ETAPA,
            "A padronização documental foi bloqueada "
            f"({diagnostico.get('status')}). O documento único foi mantido sem o "
            "acabamento automático, e o conteúdo permanece o dos componentes.",
        )
    else:
        registro.informacao(
            ETAPA,
            f"Padronização aplicada (perfil 'aviso'): "
            f"{len(diagnostico.get('correcoes_aplicadas', []))} correção(ões); "
            f"status {diagnostico.get('status')}.",
        )
    return {
        "gravado": gravado,
        "status": diagnostico.get("status"),
        "correcoes_aplicadas": diagnostico.get("correcoes_aplicadas", []),
        "correcoes_nao_aplicadas": diagnostico.get("correcoes_nao_aplicadas", []),
        "validacao_conteudo": diagnostico.get("validacao_conteudo", {}),
        "testes": diagnostico.get("testes", {}),
    }


def _pasta_de_saida(pasta_processo: Optional[Path], saida: Optional[Path],
                    manifesto: Optional[Path]) -> Path:
    if saida is not None:
        caminho = Path(saida)
        return caminho.parent if caminho.name == "saida" else caminho
    if pasta_processo is not None:
        return Path(pasta_processo) / PASTA_SAIDA_PADRAO
    return Path(manifesto).parent / PASTA_SAIDA_PADRAO


def _dados_relatorio(registro: Registro, manifesto: dict[str, Any],
                     componentes, dados_tr, plano, produzidos: dict[str, Path],
                     extras: dict[str, Any], permitir_rascunho: bool,
                     relatorio_em: Optional[Path] = None) -> dict[str, Any]:
    """Consolida tudo o que o relatório do item 17 do escopo precisa mostrar."""
    processo = (manifesto or {}).get("processo", {})
    instrumento = getattr(componentes, "instrumento", None)
    formatacao = extras.get("formatacao", {}) or {}
    preservacao = extras.get("preservacao", {}) or {}
    validacao_final = extras.get("validacao_final", {}) or {}
    pendencias = extras.get("pendencias", {}) or {}
    pacote = extras.get("pacote")

    status = relatorio_aviso_completo.status_permitido(
        decidir_status(registro, rascunho_solicitado=permitir_rascunho))

    arquivos = [str(caminho) for caminho in produzidos.values()]
    if extras.get("documento_final"):
        arquivos.append(str(extras["documento_final"]))
    if pacote:
        arquivos += [v for v in (pacote.get("documento_unico_pdf"),
                                 pacote.get("zip")) if v]
        arquivos += pacote.get("componentes_docx", [])
        arquivos += pacote.get("componentes_pdf", [])

    dados = {
        "gerado_em": date.today().isoformat(),
        "modo": "montagem" + (" (rascunho autorizado)" if permitir_rascunho else ""),
        "status": status,
        "identificacao": {
            "Processo": processo.get("numero"),
            "Dispensa": processo.get("dispensa"),
            "Objeto": processo.get("objeto"),
            "Critério de julgamento": processo.get("criterio_julgamento"),
            "Fundamento legal": processo.get("fundamento_legal"),
        },
        "componentes": [
            f"{chave}: {Path(caminho).name}" for chave, caminho in produzidos.items()
        ],
        "minutas": getattr(componentes, "minutas_utilizadas", []),
        "termo_referencia": (
            {
                "Arquivo": str(dados_tr.caminho),
                "Processo identificado no TR": dados_tr.numero_processo,
                "Objeto": dados_tr.objeto,
                "Itens": len(dados_tr.itens),
            } if dados_tr else None
        ),
        "instrumento": (
            {
                "Tipo": instrumento.tipo,
                "Minuta anexada ao aviso": "sim" if instrumento.incluir_minuta else "não",
                "Minuta selecionada": (str(instrumento.minuta)
                                       if instrumento.minuta else None),
                "Fonte da decisão": instrumento.fonte_da_decisao,
            } if instrumento else None
        ),
        "anexos": plano.como_lista() if plano else [],
        "itens_tr": [item.como_dicionario() for item in dados_tr.itens] if dados_tr else [],
        "habilitacao": [
            str(o) for o in registro.itens
            if "habilita" in o.mensagem.lower() or "DIVERG" in o.mensagem
        ],
        "proposta": [str(o) for o in registro.itens if o.etapa == "modelo de proposta"],
        "contrato": [str(o) for o in registro.itens if o.etapa == "minuta de contrato"],
        "declaracao": [
            str(o) for o in registro.itens
            if "declara" in o.mensagem.lower() and o.etapa != "modelo de proposta"
        ],
        # Só o que a validação cruzada reconheceu como pendência de verdade. A
        # varredura crua do módulo de padronização conta também as lacunas
        # reservadas ao proponente; listá-las aqui faria um documento correto
        # parecer incompleto.
        "campos_pendentes": validacao_final.get("campos_pendentes", []),
        "correcoes": (extras.get("correcoes") or [])
                     + formatacao.get("correcoes_aplicadas", []),
        "validacao_conteudo": {
            "Linhas conferidas": preservacao.get("linhas_conferidas"),
            "Linhas perdidas": preservacao.get("linhas_perdidas"),
            "Blocos 'OU' remanescentes": validacao_final.get("blocos_ou"),
            "Lacunas reservadas ao proponente": len(
                pendencias.get("campos_pendentes") or []),
            "Preservação na padronização": (
                formatacao.get("validacao_conteudo", {}).get("conteudo_preservado")),
        },
        "validacao_formatacao": {
            "Perfil aplicado": "aviso",
            "Status da padronização": formatacao.get("status", "não executada"),
            "Idempotência": formatacao.get("testes", {}).get("idempotencia"),
            "Validação visual": formatacao.get("testes", {}).get("validacao_visual"),
            "Timbre preservado": validacao_final.get("timbre", {}).get("intacto"),
        },
        "arquivos": arquivos,
        # Caminho do DOCX único, para quem precisa registrá-lo no processo
        # (Gestão Documental, item 26 do escopo daquele módulo).
        "documento_final": (str(extras["documento_final"])
                            if extras.get("documento_final") else None),
        "resumo_resultado": _resumo(registro, status),
        "ocorrencias": registro.como_lista(),
        "contagem": registro.contagem(),
    }

    if relatorio_em is not None:
        md, js = relatorio_aviso_completo.gravar(
            dados, Path(relatorio_em) / "VALIDACAO_AVISO_COMPLETO.md")
        dados["relatorio_markdown"] = str(md)
        dados["relatorio_json"] = str(js)
    return dados


def _resumo(registro: Registro, status: str) -> list[str]:
    contagem = registro.contagem()
    return [
        f"Erros bloqueantes: {contagem['ERRO BLOQUEANTE']}",
        f"Alertas: {contagem['ALERTA']}",
        f"Pendências humanas: {contagem['PENDÊNCIA HUMANA']}",
        f"Status atribuído: {status}",
    ]


# --------------------------------------------------------------------------- #
# Linha de comando
# --------------------------------------------------------------------------- #

def construir_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Monta o Aviso de Dispensa Completo da Câmara de Itanhandu.")
    parser.add_argument("--processo", type=Path,
                        help="pasta do processo em 08_processos_em_andamento/")
    parser.add_argument("--manifesto", type=Path,
                        help="manifesto_aviso_completo.json")
    parser.add_argument("--saida", type=Path,
                        help="pasta 07_AVISO_COMPLETO/ (ou .../saida/)")
    parser.add_argument("--aviso", type=Path, help="minuta-mãe do aviso")
    parser.add_argument("--tr", type=Path, help="Termo de Referência final")
    parser.add_argument("--contrato", type=Path,
                        help="minuta de contrato oficial a anexar")
    parser.add_argument("--incluir-contrato", dest="incluir_contrato",
                        action="store_true", default=None,
                        help="determinação expressa: haverá Termo de Contrato")
    parser.add_argument("--sem-contrato", dest="incluir_contrato",
                        action="store_false",
                        help="determinação expressa: instrumento equivalente")
    parser.add_argument("--autorizar-tr-rascunho", action="store_true",
                        help="aceita TR com nome de rascunho (autorização expressa)")
    parser.add_argument("--rascunho-com-pendencias", action="store_true",
                        help=f"gera assim mesmo, marcado como '{AVISO_RASCUNHO}'")
    parser.add_argument("--sem-padronizacao", action="store_true",
                        help="não aplica o módulo de padronização documental")
    parser.add_argument("--sem-pacote", action="store_true",
                        help="não gera anexos separados, PDF nem ZIP")
    parser.add_argument("--sem-pdf", action="store_true",
                        help="gera o pacote sem tentar converter para PDF")
    parser.add_argument("--registrar-em-processo", default=None, metavar="PROCESSO",
                        help="registra o AVISO_COMPLETO na Gestão Documental "
                             "(scripts/gestao_documental) como versão vigente")
    parser.add_argument("--motivo-registro", default=None,
                        help="motivo da nova versão no manifesto do processo")
    return parser


def main(argv: Optional[list[str]] = None) -> int:
    args = construir_parser().parse_args(argv)
    if not args.processo and not args.manifesto:
        print("Informe --processo ou --manifesto.", file=sys.stderr)
        return 2

    dados = montar(
        args.processo, args.manifesto, saida=args.saida, aviso=args.aviso,
        tr=args.tr, contrato=args.contrato,
        decisao_contrato=args.incluir_contrato,
        autorizar_tr_rascunho=args.autorizar_tr_rascunho,
        permitir_rascunho=args.rascunho_com_pendencias,
        padronizar=not args.sem_padronizacao,
        gerar_pacote=not args.sem_pacote,
        converter_pdf=not args.sem_pdf,
    )

    print(f"Status: {dados['status']}")
    for ocorrencia in dados["ocorrencias"]:
        if ocorrencia["severidade"] != "INFORMAÇÃO":
            print(f"  [{ocorrencia['severidade']}] {ocorrencia['etapa']}: "
                  f"{ocorrencia['mensagem']}")
    if dados.get("relatorio_markdown"):
        print(f"\nRelatório: {dados['relatorio_markdown']}")

    if args.registrar_em_processo:
        codigo = _registrar_em_processo(args, dados)
        if codigo:
            return codigo
    return 1 if dados["contagem"].get("ERRO BLOQUEANTE") else 0


def _situacao_formatacao(status: Any) -> str:
    """Traduz o status da padronização para o vocabulário do manifesto."""
    texto = str(status or "").upper()
    if not texto or "NÃO EXECUTADA" in texto or "NAO EXECUTADA" in texto:
        return "nao_executada"
    if "BLOQUE" in texto or "ERRO" in texto:
        return "reprovado"
    if "RESSALVA" in texto or "ALERTA" in texto:
        return "aprovado_com_ressalvas"
    return "aprovado"


def _registrar_em_processo(args, dados: dict[str, Any]) -> int:
    """
    Entrega o documento único à Gestão Documental (item 26 daquele módulo).

    Montagem com erro bloqueante não é registrada: um aviso que não passou na
    validação não pode virar a versão vigente do processo. Os componentes
    (TR, modelo de proposta, contrato) NÃO são registrados aqui — eles já têm
    registro próprio, e duplicá-los na raiz do processo é justamente o que a
    Gestão Documental existe para evitar.
    """
    if dados["contagem"].get("ERRO BLOQUEANTE"):
        print("[AVISO] Montagem com erro bloqueante: nada foi registrado no processo.",
              file=sys.stderr)
        return 1
    documento = dados.get("documento_final")
    if not documento:
        print("[AVISO] Não há documento único gravado para registrar.", file=sys.stderr)
        return 1

    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "gestao_documental"))
    from registrar_documento import registrar_saida_gerada  # noqa: E402

    try:
        resultado = registrar_saida_gerada(
            args.registrar_em_processo,
            "AVISO_COMPLETO",
            documento,
            minuta_origem="05_minutas/AVISO/AVISO_CONTRATACAO_DIRETA_MINUTA_MAE.docx",
            motivo=args.motivo_registro or "montagem do Aviso de Dispensa Completo",
            validacao={
                "conteudo": ("aprovado"
                             if dados["validacao_conteudo"].get("Preservação na padronização")
                             else "nao_executada"),
                "formatacao": _situacao_formatacao(
                    dados["validacao_formatacao"].get("Status da padronização")),
                "status_padronizacao": dados["validacao_formatacao"].get(
                    "Status da padronização"),
            },
        )
    except Exception as exc:  # noqa: BLE001 - a mensagem é o produto do CLI
        print(f"[ERRO] Registro no processo recusado: {exc}", file=sys.stderr)
        return 1

    print()
    print(resultado.texto())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
