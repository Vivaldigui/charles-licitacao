#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
util_ooxml.py — Utilidades OOXML de baixo nível do Módulo de Padronização Documental.

Concentra toda manipulação direta de XML do pacote DOCX. As demais partes do
módulo usam apenas as funções daqui, para que o acesso ao OOXML fique
encapsulado e testável (exigência do item 20 do escopo do módulo).

Regra permanente: cabeçalho, rodapé e mídia (brasão/timbre) são partes
PROTEGIDAS. Nenhuma função deste módulo escreve em `word/header*.xml`,
`word/footer*.xml` ou `word/media/*`.
"""
from __future__ import annotations

import hashlib
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Iterable, Iterator, Optional

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def qn(tag: str) -> str:
    """Converte 'w:tag' na forma Clark '{ns}tag'."""
    prefixo, _, local = tag.partition(":")
    if prefixo != "w":
        raise ValueError(f"Prefixo não suportado: {prefixo}")
    return f"{{{W_NS}}}{local}"


# --------------------------------------------------------------------------- #
# Dependência opcional
# --------------------------------------------------------------------------- #

MSG_SEM_DOCX = (
    "A biblioteca python-docx não está instalada. O Módulo de Padronização "
    "Documental depende dela.\n"
    "Instale com:  python -m pip install -r requirements-docx.txt"
)


def exigir_python_docx():
    """Importa python-docx ou levanta erro com instrução de instalação."""
    try:
        import docx  # noqa: F401
    except ImportError as exc:  # pragma: no cover - depende do ambiente
        raise RuntimeError(MSG_SEM_DOCX) from exc
    return docx


# --------------------------------------------------------------------------- #
# Marcadores de campo
# --------------------------------------------------------------------------- #

CAMPO_RE = re.compile(r"\{\{([A-Z0-9_]+)\}\}")
PREENCHER_RE = re.compile(r"\[PREENCHER:?[^\]]*\]", re.IGNORECASE)
DEFINIR_RE = re.compile(r"\((?:\s*)definir[^)]*\)", re.IGNORECASE)
SUBLINHADO_RE = re.compile(r"_{4,}")
CHAVE_RESIDUAL_RE = re.compile(r"\{\{|\}\}")
BLOCO_OU_RE = re.compile(r"^\s*OU\s*$", re.IGNORECASE)
OPCAO_NAO_MARCADA_RE = re.compile(r"\(\s*\)")


# --------------------------------------------------------------------------- #
# Iteração de parágrafos e tabelas
# --------------------------------------------------------------------------- #

def iter_paragrafos_tabela(tabela) -> Iterator:
    """Percorre parágrafos de uma tabela, inclusive tabelas aninhadas."""
    for linha in tabela.rows:
        for celula in linha.cells:
            for paragrafo in celula.paragraphs:
                yield paragrafo
            for sub in celula.tables:
                yield from iter_paragrafos_tabela(sub)


def iter_paragrafos_corpo(documento) -> Iterator:
    """Todos os parágrafos do corpo, incluindo os de dentro de tabelas."""
    for paragrafo in documento.paragraphs:
        yield paragrafo
    for tabela in documento.tables:
        yield from iter_paragrafos_tabela(tabela)


def iter_tabelas(documento) -> Iterator:
    """Tabelas de primeiro nível e aninhadas."""
    def _desce(tabelas):
        for tabela in tabelas:
            yield tabela
            for linha in tabela.rows:
                for celula in linha.cells:
                    yield from _desce(celula.tables)
    yield from _desce(documento.tables)


def partes_cabecalho_rodape(secao) -> list:
    """
    Cabeçalhos/rodapés que já EXISTEM na seção.

    Cuidado deliberado: em python-docx, ler `first_page_header` ou
    `even_page_header` de uma seção que não os define CRIA a parte no pacote.
    Isso alteraria o timbre só por auditar. Por isso as variantes só são lidas
    quando o documento declara usá-las, e partes vinculadas à seção anterior
    (`is_linked_to_previous`) são puladas.
    """
    partes = [secao.header, secao.footer]
    if secao.different_first_page_header_footer:
        partes += [secao.first_page_header, secao.first_page_footer]
    return [p for p in partes if p is not None and not p.is_linked_to_previous]


def iter_paragrafos_cabecalho_rodape(documento) -> Iterator:
    """Parágrafos de cabeçalhos e rodapés — LEITURA apenas (auditoria)."""
    for secao in documento.sections:
        for parte in partes_cabecalho_rodape(secao):
            for paragrafo in parte.paragraphs:
                yield paragrafo
            for tabela in parte.tables:
                yield from iter_paragrafos_tabela(tabela)


def texto_paragrafo(paragrafo) -> str:
    """Texto do parágrafo lendo todos os w:t (inclusive dentro de hyperlink)."""
    return "".join(t.text or "" for t in paragrafo._p.iter(qn("w:t")))


def elementos_texto(paragrafo) -> list:
    """Lista dos elementos w:t do parágrafo, na ordem do documento."""
    return list(paragrafo._p.iter(qn("w:t")))


def runs_do_paragrafo(paragrafo) -> list:
    """Elementos w:r do parágrafo, inclusive dentro de w:hyperlink."""
    return list(paragrafo._p.iter(qn("w:r")))


# --------------------------------------------------------------------------- #
# Substituição de marcadores preservando formatação
# --------------------------------------------------------------------------- #

@dataclass
class ResultadoSubstituicao:
    """Resultado de uma passagem de substituição de campos."""

    preenchidos: list[str] = field(default_factory=list)
    nao_resolvidos: list[str] = field(default_factory=list)


def substituir_marcadores_paragrafo(
    paragrafo,
    resolver: Callable[[str], Optional[str]],
    padrao: re.Pattern[str] = CAMPO_RE,
) -> ResultadoSubstituicao:
    """
    Substitui marcadores `{{CAMPO}}` mesmo quando divididos entre vários runs.

    A formatação-base preservada é a do run onde o marcador COMEÇA — é ela que
    carrega o estilo do campo na minuta. Nenhum run novo é criado, de modo que
    o texto inserido não vira um fragmento de fonte diferente.

    `resolver(nome)` devolve o texto de substituição ou `None` para deixar o
    marcador intacto.
    """
    resultado = ResultadoSubstituicao()
    elementos = elementos_texto(paragrafo)
    if not elementos:
        return resultado

    textos = [e.text or "" for e in elementos]
    completo = "".join(textos)
    if not padrao.search(completo):
        return resultado

    # Dono de cada caractere: índice do w:t de origem.
    dono: list[int] = []
    for indice, trecho in enumerate(textos):
        dono.extend([indice] * len(trecho))

    saida: list[str] = []
    saida_dono: list[int] = []
    posicao = 0
    houve_troca = False

    for casamento in padrao.finditer(completo):
        nome = casamento.group(1) if casamento.groups() else casamento.group(0)
        valor = resolver(nome)
        if valor is None:
            resultado.nao_resolvidos.append(nome)
            continue
        inicio, fim = casamento.span()
        saida.append(completo[posicao:inicio])
        saida_dono.extend(dono[posicao:inicio])
        # Todo o texto inserido herda o run onde o marcador começou.
        ancora = dono[inicio] if inicio < len(dono) else len(textos) - 1
        saida.append(valor)
        saida_dono.extend([ancora] * len(valor))
        posicao = fim
        houve_troca = True
        resultado.preenchidos.append(nome)

    if not houve_troca:
        return resultado

    saida.append(completo[posicao:])
    saida_dono.extend(dono[posicao:])

    novos = [""] * len(textos)
    for caractere, indice in zip("".join(saida), saida_dono):
        novos[indice] += caractere

    for elemento, novo in zip(elementos, novos):
        definir_texto(elemento, novo)

    return resultado


def definir_texto(elemento_t, texto: str) -> None:
    """Grava texto em um w:t preservando espaços significativos."""
    elemento_t.text = texto
    chave = "{http://www.w3.org/XML/1998/namespace}space"
    if texto != texto.strip() or texto == "":
        elemento_t.set(chave, "preserve")
    elif chave in elemento_t.attrib:
        del elemento_t.attrib[chave]


def marcadores_pendentes(texto: str) -> list[str]:
    """Marcadores de preenchimento ainda presentes em um texto."""
    achados: list[str] = []
    achados += [f"{{{{{n}}}}}" for n in CAMPO_RE.findall(texto)]
    achados += PREENCHER_RE.findall(texto)
    achados += DEFINIR_RE.findall(texto)
    achados += SUBLINHADO_RE.findall(texto)
    if not achados and CHAVE_RESIDUAL_RE.search(texto):
        achados.append("chaves residuais ({{ ou }})")
    return achados


# --------------------------------------------------------------------------- #
# Propriedades de parágrafo (paginação)
# --------------------------------------------------------------------------- #

def _pPr(paragrafo):
    return paragrafo._p.get_or_add_pPr()


def _flag(paragrafo, tag: str, ativo: bool) -> bool:
    """Liga/desliga um flag booleano de w:pPr. Devolve True se mudou algo."""
    pPr = _pPr(paragrafo)
    existente = pPr.find(qn(tag))
    if ativo:
        if existente is not None and existente.get(qn("w:val")) in (None, "1", "true", "on"):
            return False
        if existente is None:
            existente = pPr.makeelement(qn(tag), {})
            pPr.append(existente)
        if qn("w:val") in existente.attrib:
            del existente.attrib[qn("w:val")]
        return True
    if existente is None:
        return False
    pPr.remove(existente)
    return True


def definir_manter_com_proximo(paragrafo, ativo: bool = True) -> bool:
    """keep_with_next — impede título isolado no fim da página."""
    return _flag(paragrafo, "w:keepNext", ativo)


def definir_manter_junto(paragrafo, ativo: bool = True) -> bool:
    """keep_together — impede quebra no meio do parágrafo."""
    return _flag(paragrafo, "w:keepLines", ativo)


def definir_controle_viuvas(paragrafo, ativo: bool = True) -> bool:
    """widow_control — controle de linhas órfãs e viúvas."""
    return _flag(paragrafo, "w:widowControl", ativo)


def controle_viuvas_desligado(paragrafo) -> bool:
    """True quando o parágrafo desliga explicitamente o controle de viúvas."""
    pPr = paragrafo._p.find(qn("w:pPr"))
    if pPr is None:
        return False
    elemento = pPr.find(qn("w:widowControl"))
    if elemento is None:
        return False
    return elemento.get(qn("w:val")) in ("0", "false", "off")


def definir_quebra_antes(paragrafo, ativo: bool) -> bool:
    """page_break_before."""
    return _flag(paragrafo, "w:pageBreakBefore", ativo)


def quebras_de_pagina(paragrafo) -> int:
    """Quantidade de quebras manuais de página dentro do parágrafo."""
    total = 0
    for br in paragrafo._p.iter(qn("w:br")):
        if br.get(qn("w:type")) == "page":
            total += 1
    return total


def remover_quebras_de_pagina(paragrafo) -> int:
    """Remove quebras manuais de página do parágrafo. Devolve quantas saíram."""
    removidas = 0
    for br in list(paragrafo._p.iter(qn("w:br"))):
        if br.get(qn("w:type")) == "page":
            pai = br.getparent()
            if pai is not None:
                pai.remove(br)
                removidas += 1
    return removidas


def remover_paragrafo(paragrafo) -> bool:
    """Remove o parágrafo do documento."""
    elemento = paragrafo._p
    pai = elemento.getparent()
    if pai is None:
        return False
    pai.remove(elemento)
    return True


def paragrafo_vazio(paragrafo) -> bool:
    """
    Parágrafo sem texto, sem imagem, sem quebra de página e sem quebra de seção.

    A quebra de seção conta: um `w:sectPr` dentro do `w:pPr` carrega orientação,
    margens e as referências de cabeçalho e rodapé daquele trecho. O parágrafo
    que o hospeda costuma ser visualmente vazio — e removê-lo como "linha em
    branco" apagaria o timbre e a configuração de página junto.
    """
    if texto_paragrafo(paragrafo).strip():
        return False
    if next(paragrafo._p.iter(qn("w:drawing")), None) is not None:
        return False
    if next(paragrafo._p.iter(qn("w:pict")), None) is not None:
        return False
    if quebras_de_pagina(paragrafo):
        return False
    if next(paragrafo._p.iter(qn("w:sectPr")), None) is not None:
        return False
    return True


# --------------------------------------------------------------------------- #
# Tabelas
# --------------------------------------------------------------------------- #

def definir_repetir_cabecalho(linha, ativo: bool = True) -> bool:
    """w:tblHeader — repete a linha de cabeçalho nas páginas seguintes."""
    trPr = linha._tr.get_or_add_trPr()
    existente = trPr.find(qn("w:tblHeader"))
    if ativo:
        if existente is not None:
            return False
        trPr.append(trPr.makeelement(qn("w:tblHeader"), {}))
        return True
    if existente is None:
        return False
    trPr.remove(existente)
    return True


def definir_nao_dividir_linha(linha, ativo: bool = True) -> bool:
    """w:cantSplit — impede que a linha se parta entre páginas."""
    trPr = linha._tr.get_or_add_trPr()
    existente = trPr.find(qn("w:cantSplit"))
    if ativo:
        if existente is not None:
            return False
        trPr.append(trPr.makeelement(qn("w:cantSplit"), {}))
        return True
    if existente is None:
        return False
    trPr.remove(existente)
    return True


def sombrear_celula(celula, cor_hex: str) -> bool:
    """Aplica sombreamento sólido em uma célula. Devolve True se mudou."""
    tcPr = celula._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = tcPr.makeelement(qn("w:shd"), {})
        tcPr.append(shd)
    mudou = (
        shd.get(qn("w:fill")) != cor_hex
        or shd.get(qn("w:val")) != "clear"
    )
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor_hex)
    return mudou


def definir_alinhamento_vertical(celula, valor: str = "center") -> bool:
    """w:vAlign da célula."""
    tcPr = celula._tc.get_or_add_tcPr()
    valign = tcPr.find(qn("w:vAlign"))
    if valign is None:
        valign = tcPr.makeelement(qn("w:vAlign"), {})
        tcPr.append(valign)
    if valign.get(qn("w:val")) == valor:
        return False
    valign.set(qn("w:val"), valor)
    return True


def definir_margens_internas_tabela(tabela, cm: float) -> bool:
    """w:tblCellMar uniforme (em centímetros)."""
    twips = str(int(round(cm * 567)))
    tblPr = tabela._tbl.tblPr
    mar = tblPr.find(qn("w:tblCellMar"))
    if mar is None:
        mar = tblPr.makeelement(qn("w:tblCellMar"), {})
        tblPr.append(mar)
    mudou = False
    for lado in ("top", "left", "bottom", "right"):
        elemento = mar.find(qn(f"w:{lado}"))
        if elemento is None:
            elemento = mar.makeelement(qn(f"w:{lado}"), {})
            mar.append(elemento)
        if elemento.get(qn("w:w")) != twips or elemento.get(qn("w:type")) != "dxa":
            mudou = True
        elemento.set(qn("w:w"), twips)
        elemento.set(qn("w:type"), "dxa")
    return mudou


def celulas_mescladas(tabela) -> int:
    """Conta células com mesclagem horizontal ou vertical."""
    total = 0
    for tc in tabela._tbl.iter(qn("w:tc")):
        tcPr = tc.find(qn("w:tcPr"))
        if tcPr is None:
            continue
        if tcPr.find(qn("w:gridSpan")) is not None or tcPr.find(qn("w:vMerge")) is not None:
            total += 1
    return total


# --------------------------------------------------------------------------- #
# Partes protegidas e metadados do pacote
# --------------------------------------------------------------------------- #

PARTES_PROTEGIDAS = re.compile(r"^word/(header\d*\.xml|footer\d*\.xml)$")
RE_ALVO_MIDIA = re.compile(r'Target="([^"]*media/[^"]+)"')


def midia_referenciada(caminho: Path) -> set[str]:
    """
    Mídia efetivamente usada, segundo os arquivos `.rels` do pacote.

    Distingue o brasão/timbre (referenciado por `header*.xml.rels` e
    `footer*.xml.rels`) das imagens órfãs que sobram de edições anteriores e que
    o Word e o python-docx descartam ao salvar. Só a mídia REFERENCIADA é parte
    protegida — perder órfã é limpeza, perder referenciada é destruir o timbre.
    """
    referenciada: set[str] = set()
    with zipfile.ZipFile(caminho, "r") as zf:
        for nome in zf.namelist():
            if not nome.endswith(".rels"):
                continue
            xml = zf.read(nome).decode("utf-8", "ignore")
            for alvo in RE_ALVO_MIDIA.findall(xml):
                arquivo = alvo.split("media/")[-1]
                referenciada.add(f"word/media/{arquivo}")
    return referenciada


def midia_orfa(caminho: Path) -> set[str]:
    """Imagens presentes no pacote que nenhum `.rels` referencia."""
    todas = {n for n in partes_do_pacote(caminho) if n.startswith("word/media/")}
    return todas - midia_referenciada(caminho)


RE_T_XML = re.compile(rb"<w:t[^>]*>(.*?)</w:t>", re.S)
RE_EMBED = re.compile(rb'r:embed="([^"]+)"')
RE_EXTENT = re.compile(rb'<wp:extent\s+cx="(\d+)"\s+cy="(\d+)"')


def assinatura_semantica_cabecalho(xml: bytes) -> str:
    """
    Assinatura do que importa em um cabeçalho/rodapé.

    Não usa o hash bruto do XML: qualquer biblioteca que abra e salve o DOCX
    reserializa a parte e muda os bytes sem mudar nada visível. A assinatura
    considera o que de fato caracteriza o timbre — textos, quantidade de
    imagens, quais imagens são referenciadas e em que dimensões. Assim, texto
    trocado, brasão removido ou imagem redimensionada acusam divergência, e a
    reserialização inocente não acusa.
    """
    textos = b"\x1f".join(t.strip() for t in RE_T_XML.findall(xml))
    desenhos = xml.count(b"<w:drawing>")
    embeds = b",".join(sorted(set(RE_EMBED.findall(xml))))
    extents = b",".join(sorted(b"x".join(par) for par in RE_EXTENT.findall(xml)))
    material = b"|".join([textos, str(desenhos).encode(), embeds, extents])
    return hashlib.sha256(material).hexdigest()


def assinaturas_partes(caminho: Path) -> dict[str, str]:
    """
    Assinatura de cada parte protegida: cabeçalhos, rodapés e mídia referenciada.

    Cabeçalho/rodapé usam assinatura semântica; as imagens usam hash de bytes,
    porque nelas qualquer alteração binária é perda real. Mídia órfã fica de
    fora de propósito — ver `midia_referenciada`.
    """
    caminho = Path(caminho)
    referenciada = midia_referenciada(caminho)
    assinaturas: dict[str, str] = {}
    with zipfile.ZipFile(caminho, "r") as zf:
        for nome in sorted(zf.namelist()):
            if PARTES_PROTEGIDAS.match(nome):
                assinaturas[nome] = assinatura_semantica_cabecalho(zf.read(nome))
            elif nome in referenciada:
                assinaturas[nome] = hashlib.sha256(zf.read(nome)).hexdigest()
    return assinaturas


def partes_do_pacote(caminho: Path) -> set[str]:
    with zipfile.ZipFile(caminho, "r") as zf:
        return set(zf.namelist())


def possui_parte(caminho: Path, nome: str) -> bool:
    return nome in partes_do_pacote(caminho)


def contagem_controle_alteracoes(caminho: Path) -> int:
    """Número de marcas de controle de alterações em word/document.xml."""
    with zipfile.ZipFile(caminho, "r") as zf:
        xml = zf.read("word/document.xml").decode("utf-8", "ignore")
    return len(re.findall(r"<w:(ins|del|moveFrom|moveTo)\b", xml))


def contagem_texto_oculto(caminho: Path) -> int:
    with zipfile.ZipFile(caminho, "r") as zf:
        xml = zf.read("word/document.xml").decode("utf-8", "ignore")
    return len(re.findall(r"<w:vanish\b(?![^>]*w:val=\"(?:0|false|off)\")", xml))


def propriedades_pessoais(caminho: Path) -> dict[str, str]:
    """Autor, último editor e empresa registrados no pacote."""
    dados: dict[str, str] = {}
    with zipfile.ZipFile(caminho, "r") as zf:
        for parte, campos in (
            ("docProps/core.xml", ("creator", "lastModifiedBy")),
            ("docProps/app.xml", ("Company", "Manager")),
        ):
            if parte not in zf.namelist():
                continue
            xml = zf.read(parte).decode("utf-8", "ignore")
            for campo in campos:
                achado = re.search(rf"<(?:\w+:)?{campo}>([^<]*)</(?:\w+:)?{campo}>", xml)
                if achado and achado.group(1).strip():
                    dados[campo] = achado.group(1).strip()
    return dados


# --------------------------------------------------------------------------- #
# Normalização de conteúdo (validação de preservação)
# --------------------------------------------------------------------------- #

def normalizar_texto(texto: str) -> str:
    """Texto comparável: NFC, espaços colapsados, sem espaços nas pontas."""
    texto = unicodedata.normalize("NFC", texto)
    texto = texto.replace(" ", " ").replace("\t", " ")
    return re.sub(r"\s+", " ", texto).strip()


def hash_conteudo(linhas: Iterable[str]) -> str:
    """Assinatura estável do conteúdo normalizado."""
    digest = hashlib.sha256()
    for linha in linhas:
        digest.update(linha.encode("utf-8"))
        digest.update(b"\n")
    return digest.hexdigest()
